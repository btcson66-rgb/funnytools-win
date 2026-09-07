"""services/pdf_table.py — 座標推欄 fallback 與 XLSX 輸出完整性。"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from services.pdf_table import (
    _cluster_rows,
    _table_from_words,
    extract_tables,
    pdf_tables_to_xlsx,
    tables_to_xlsx,
)


def word(text: str, x0: float, top: float, width: float = 40, height: float = 10) -> dict:
    return {"text": text, "x0": x0, "x1": x0 + width, "top": top, "bottom": top + height}


def grid(rows: int, cols: int, col_gap: float = 120, row_gap: float = 20) -> list[dict]:
    return [
        word(f"r{r}c{c}", 50 + c * col_gap, 100 + r * row_gap)
        for r in range(rows)
        for c in range(cols)
    ]


# ------------------------------------------------------------------ _cluster_rows


def test_cluster_rows_groups_by_vertical_proximity_and_sorts_left_to_right():
    words = [word("b", 200, 100), word("a", 50, 102), word("c", 50, 300)]
    rows = _cluster_rows(words)
    assert [[w["text"] for w in row] for row in rows] == [["a", "b"], ["c"]]


def test_cluster_rows_keeps_rows_apart_beyond_the_tolerance():
    rows = _cluster_rows([word("a", 50, 100), word("b", 50, 106)])
    assert len(rows) == 2  # 6pt > 5.0 tolerance


def test_cluster_rows_on_empty_input():
    assert _cluster_rows([]) == []


# -------------------------------------------------------------- _table_from_words


def test_table_from_words_builds_a_grid_from_aligned_columns():
    table = _table_from_words(grid(rows=4, cols=3))
    assert len(table) == 4
    assert all(len(row) == 3 for row in table)
    assert table[0] == ["r0c0", "r0c1", "r0c2"]


@pytest.mark.parametrize("count", [0, 1, 2, 3])
def test_table_from_words_refuses_fewer_than_four_words(count):
    """守門條件 len(words) < 4，避免把零星文字當成表格。"""
    assert _table_from_words(grid(rows=1, cols=1)[:count]) == []


def test_table_from_words_refuses_two_row_layouts():
    """兩行對齊文字太容易誤判成表格，程式碼刻意要求至少 3 列。"""
    assert _table_from_words(grid(rows=2, cols=3)) == []


def test_table_from_words_refuses_a_single_column():
    assert _table_from_words(grid(rows=5, cols=1)) == []


def test_table_from_words_false_positives_on_incidentally_aligned_prose_known_gap():
    """已知缺口（資料品質，非安全）：三行散文只要左緣「碰巧」在 12pt 內對齊，
    就會被推成表格。

    程式碼註解宣稱「normal prose usually does not」有重複的左緣錨點，但
    x_tolerance = max(12.0, 字高*0.65)，一般字高 10pt 時就是固定 12pt，
    而 min_occurrence 只要求 3 列中出現 2 列。下面這段完全是散文，
    卻被推成 3 欄表格。此測試鎖住現況，若日後把 fallback 收緊，這裡會紅。
    """
    words = [
        word("Lorem", 50, 100), word("ipsum", 95, 100), word("dolor", 143, 100),
        word("sit", 72, 120), word("amet", 96, 120), word("consectetur", 137, 120),
        word("adipiscing", 61, 140), word("elit", 130, 140), word("sed", 158, 140),
    ]
    result = _table_from_words(words)
    assert result != [], "現況就是會誤判；若已修好請一併更新這個測試"
    assert result[0] == ["Lorem", "ipsum", "dolor"]


def test_table_from_words_ignores_rows_with_only_one_populated_cell():
    words = grid(rows=4, cols=3)
    words.append(word("orphan", 50, 500))  # 自成一列且只有一個值
    table = _table_from_words(words)
    assert ["orphan", "", ""] not in table


def test_table_from_words_tolerates_zero_height_words():
    """height 全為 0 時 median 會走 fallback 12.0，不能除以零或炸掉。"""
    words = [
        word(f"r{r}c{c}", 50 + c * 120, 100 + r * 20, height=0)
        for r in range(4)
        for c in range(2)
    ]
    assert len(_table_from_words(words)) == 4


# ----------------------------------------------------------------- tables_to_xlsx


def sheet_of(xlsx: bytes):
    return load_workbook(io.BytesIO(xlsx))


def test_tables_to_xlsx_empty_input_produces_a_placeholder_sheet():
    wb = sheet_of(tables_to_xlsx([]))
    assert wb.sheetnames == ["No tables"]
    assert wb["No tables"]["A1"].value == "No table detected"


def test_tables_to_xlsx_writes_one_sheet_per_table_with_page_table_naming():
    xlsx = tables_to_xlsx([
        {"page": 1, "table": 1, "rows": [["a", "b"], ["c", "d"]]},
        {"page": 2, "table": 3, "rows": [["e"]]},
    ])
    wb = sheet_of(xlsx)
    assert wb.sheetnames == ["P1_T1", "P2_T3"]
    assert [c.value for c in wb["P1_T1"][1]] == ["a", "b"]


def test_tables_to_xlsx_disambiguates_colliding_sheet_names():
    xlsx = tables_to_xlsx([{"page": 1, "table": 1, "rows": [["a"]]}] * 3)
    names = sheet_of(xlsx).sheetnames
    assert len(names) == len(set(names)) == 3


def test_tables_to_xlsx_sheet_names_stay_within_excel_31_char_limit():
    xlsx = tables_to_xlsx([
        {"page": 9_999_999_999, "table": 9_999_999_999, "rows": [["a"]]},
        {"page": 9_999_999_999, "table": 9_999_999_999, "rows": [["b"]]},
    ])
    assert all(len(n) <= 31 for n in sheet_of(xlsx).sheetnames)


def test_tables_to_xlsx_stringifies_every_cell_type():
    xlsx = tables_to_xlsx([
        {"page": 1, "table": 1, "rows": [[1, 2.5, None, True, "=1+1", "-cmd"]]}
    ])
    values = [c.value for c in sheet_of(xlsx)["P1_T1"][1]]
    # None 會被寫成空字串，openpyxl 讀回來是 None（空儲存格）。
    assert values == ["1", "2.5", None, "True", "=1+1", "-cmd"]


@pytest.mark.parametrize(
    "payload",
    ['=HYPERLINK("http://evil.example","click")', "=1+1", "=cmd|'/c calc'!A1"],
)
def test_tables_to_xlsx_neutralises_formula_cells(payload):
    """openpyxl 的 Cell._bind_value 會把長度 >1 且以 "=" 開頭的字串標成
    data_type='f'，也就是寫成真正的 Excel 公式。惡意 PDF 表格、或直接
    POST /api/pdf/export-tables 的 JSON，都能靠這個把公式塞進使用者的試算表。"""
    cell = sheet_of(tables_to_xlsx([{"page": 1, "table": 1, "rows": [[payload]]}]))["P1_T1"]["A1"]
    assert cell.data_type == "s", "儲存格仍是公式型別，Excel 開啟時會執行它"
    # 中和方式是改型別而不是加單引號前綴，所以讀回來的值必須完全沒變
    # （加前綴會把 ' 寫進資料本身，破壞非 Excel 讀取端看到的內容）。
    assert cell.value == payload


def test_tables_to_xlsx_leaves_non_formula_cells_as_plain_strings():
    """+ - @ tab CR 在 openpyxl 不會被轉成公式節點，中和步驟不得誤傷它們。"""
    row = sheet_of(tables_to_xlsx([
        {"page": 1, "table": 1, "rows": [["-cmd", "+1", "@SUM", "plain"]]},
    ]))["P1_T1"][1]
    assert [c.data_type for c in row] == ["s", "s", "s", "s"]
    assert [c.value for c in row] == ["-cmd", "+1", "@SUM", "plain"]


def test_tables_to_xlsx_rejects_a_non_list_row():
    with pytest.raises(ValueError, match="Each table row must be a list"):
        tables_to_xlsx([{"page": 1, "table": 1, "rows": ["not-a-list"]}])


@pytest.mark.parametrize("cell", [{"a": 1}, ["nested"]])
def test_tables_to_xlsx_rejects_non_scalar_cells(cell):
    with pytest.raises(ValueError, match="Table cells must be scalar values"):
        tables_to_xlsx([{"page": 1, "table": 1, "rows": [[cell]]}])


def test_tables_to_xlsx_rejects_a_non_numeric_page_with_a_python_message_known_gap():
    """已知缺口（低風險）：int() 的原生訊息會被原樣回給 client。"""
    with pytest.raises(ValueError, match="invalid literal for int"):
        tables_to_xlsx([{"page": "one", "table": 1, "rows": [["a"]]}])


def test_tables_to_xlsx_raises_attribute_error_when_an_entry_is_not_a_dict_known_gap():
    """已知缺口：非 dict 的 table 會拋 AttributeError，被 app 轉成 500 而非 400。"""
    with pytest.raises(AttributeError):
        tables_to_xlsx(["not-a-dict"])


def test_tables_to_xlsx_handles_ragged_rows_without_crashing():
    xlsx = tables_to_xlsx([{"page": 1, "table": 1, "rows": [["a"], ["b", "c", "d"], []]}])
    ws = sheet_of(xlsx)["P1_T1"]
    assert ws.max_column == 3


# --------------------------------------------------- extract_tables（合成 PDF）


def test_extract_tables_rejects_a_page_index_outside_the_document(table_pdf):
    with pytest.raises(ValueError, match="Requested page 5 does not exist"):
        extract_tables(table_pdf, "t.pdf", pages=[4], ocr_mode="off")


def test_extract_tables_rejects_a_negative_page_index(table_pdf):
    with pytest.raises(ValueError, match="does not exist"):
        extract_tables(table_pdf, "t.pdf", pages=[-1], ocr_mode="off")


@pytest.mark.parametrize("mode", ["Auto", "OFF", "yes", "", "1"])
def test_extract_tables_rejects_ocr_modes_off_the_allowlist(table_pdf, mode):
    with pytest.raises(ValueError, match="ocr_mode must be auto, force, or off"):
        extract_tables(table_pdf, "t.pdf", ocr_mode=mode)


def test_extract_tables_rejects_an_unsupported_ocr_language(table_pdf):
    with pytest.raises(ValueError, match="Unsupported OCR language"):
        extract_tables(table_pdf, "t.pdf", ocr_mode="off", ocr_lang="jpn")


def test_extract_tables_finds_the_coordinate_fallback_table(table_pdf):
    tables, stats = extract_tables(table_pdf, "t.pdf", ocr_mode="off")
    assert stats["pages_scanned"] == 1
    assert tables, "合成的對齊三欄文字應該被座標 fallback 抓到"
    assert tables[0]["method"] == "coordinate-fallback"
    assert tables[0]["page"] == 1
    assert any("Widget" in cell for row in tables[0]["rows"] for cell in row)


def test_extract_tables_deduplicates_and_sorts_requested_pages(table_pdf):
    _, stats = extract_tables(table_pdf, "t.pdf", pages=[0, 0, 0], ocr_mode="off")
    assert stats["pages_scanned"] == 1


def test_pdf_tables_to_xlsx_round_trips_into_a_readable_workbook(table_pdf):
    xlsx, tables, stats = pdf_tables_to_xlsx(table_pdf, "t.pdf", ocr_mode="off")
    wb = sheet_of(xlsx)
    assert stats["tables_found"] == len(tables)
    assert wb.sheetnames  # 一定至少有一個工作表（沒表格時是 "No tables"）
    assert xlsx[:2] == b"PK"  # XLSX 是 ZIP 容器
